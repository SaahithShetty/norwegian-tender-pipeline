"""Tests for tender-annex parsing.

These run against a synthetic workbook shaped like the real LIS Prisskjema, so the
suite passes without the account-walled annex present.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.parsing.annex import load_annexes, parse_annex, packs_for_atc


@pytest.fixture()
def annex(tmp_path: Path) -> Path:
    """A workbook mirroring the real annex: banner row, header on row 3, error cells."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Prisskjema"
    sheet["A2"] = "Leverandør bes fylle i gule felt"
    headers = [
        "VARENR", "VARENAVN", "LEGEMIDDELFORM", "STYRKE", "ENHET", "LEVERANDØR",
        "TILBUDT GIP med 2 desimaler", "ATC", "PAKNINGER 2021",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(3, column, header)

    rows = [
        ("103854", "Inlyta tab 1mg", "Tablett", "1mg", "56 ENPAC",
         "Pfizer Norge AS", None, "L01EK01", 179),
        ("599010", "Inlyta tab 5mg", "Tablett", "5mg", "56 ENPAC",
         "Pfizer Norge AS", None, "L01EK01", 634),
        ("44580", "Afinitor tab 5mg", "Tablett", "5mg", "30 ENPAC",
         "Novartis Norge AS", None, "L01EG02", 664),
        # Excel error literal in the volume column - must not become a number
        ("96106", "Cometriq kaps 20mg", "Kapsel", "20mg", "84 DATOP",
         "IPSEN AB", None, "L01EX07", "#N/A"),
    ]
    for index, row in enumerate(rows, start=4):
        for column, value in enumerate(row, start=1):
            sheet.cell(index, column, value)

    path = tmp_path / "LIS 2207 - Vedlegg 03 Prisskjema v 2.xlsx"
    workbook.save(path)
    return path


def test_parses_pack_level_detail(annex: Path) -> None:
    packs = parse_annex(annex)
    assert len(packs) == 4

    inlyta = next(p for p in packs if p.item_number == "103854")
    assert inlyta.product_name == "Inlyta tab 1mg"
    assert inlyta.strength == "1mg"
    assert inlyta.pack_size == "56 ENPAC"
    assert inlyta.supplier == "Pfizer Norge AS"
    assert inlyta.packs_last_12m == 179.0


def test_excel_error_literals_are_not_treated_as_data(annex: Path) -> None:
    """'#N/A' must read as absent, not as a volume or a product string."""
    cometriq = next(p for p in parse_annex(annex) if p.item_number == "96106")
    assert cometriq.packs_last_12m is None


def test_offered_price_is_empty_because_the_form_is_a_blank_template(annex: Path) -> None:
    """Suppliers fill in TILBUDT GIP when bidding, so maxPrice cannot come from here."""
    assert all(pack.offered_price is None for pack in parse_annex(annex))


def test_packs_are_selected_by_exact_atc_code(annex: Path) -> None:
    """L01EX07 now denotes cabozantinib, not axitinib.

    Exact matching keeps Cometriq out of axitinib's rows; a prefix or fuzzy match
    would attribute another manufacturer's packs to the wrong molecule.
    """
    packs = parse_annex(annex)
    axitinib = list(packs_for_atc(packs, "L01EK01"))
    assert {p.item_number for p in axitinib} == {"103854", "599010"}
    assert all("Cometriq" not in (p.product_name or "") for p in axitinib)


def test_missing_annex_directory_is_not_an_error(tmp_path: Path) -> None:
    """A reviewer without a supplier account must still get a complete run."""
    assert load_annexes(tmp_path / "does-not-exist") == []


def test_pack_rows_never_carry_the_suppliers_offered_price(annex: Path) -> None:
    """maxPrice means the buyer's regulated maximum, not a bid.

    The annex column is TILBUDT GIP - what a supplier offers when bidding. Copying it
    into maxPrice would misreport a bid as a price ceiling, so pack rows must leave
    the column empty even if a future annex arrives with that column filled in.
    """
    from src.matching.base import Match
    from src.models import DetectionMethod, SourceNotice
    from src.config import MOLECULES_BY_NAME
    from src.output import build_pack_rows
    from src.parsing.annex import parse_annex, packs_for_atc

    packs = list(packs_for_atc(parse_annex(annex), "L01EK01"))
    # Simulate an annex whose price column *is* populated.
    for pack in packs:
        pack.offered_price = 1234.56

    notice = SourceNotice(
        notice_id="2022-324116",
        title="LIS 2207 Onkologi",
        source_url="https://example.test/1",
        source_name="doffin",
    )
    match = Match(
        molecule=MOLECULES_BY_NAME["axitinib"],
        method=DetectionMethod.THERAPEUTIC_BUNDLE,
        variant="Onkologi",
        confidence=0.35,
    )

    rows = build_pack_rows(notice, match, packs)
    assert rows, "expected a row per pack"
    assert all(row.maxPrice is None for row in rows)
    # The volume, by contrast, is the buyer's own published figure and is kept.
    assert any(row.packsSoldLast12m for row in rows)
