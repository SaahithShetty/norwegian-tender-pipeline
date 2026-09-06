"""Tests for the regulated maximum-price register."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.parsing.price_register import parse_price_register


@pytest.fixture()
def register(tmp_path: Path) -> Path:
    """A workbook shaped like the published file: title row, headers on row 3."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Pakningspriser"
    headers = [
        "Varenummer", "Utgående varenummer", "Handelsnavn", "Innehaver", "Virkestoff",
        "Legemiddelform", "Styrke", "Pakningstype", "Multippel", "Antall beholdere",
        "Mengde per beholder", "Måle-enhet", "Maks AIP Gyldig", "Maks AUP Gyldig",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(3, column, header)
    sheet.append  # noqa: B018 - openpyxl rows written explicitly below
    rows = [
        ("103854", "", "Inlyta", "Pfizer AS", "Aksitinib", "Tablett", "1 mg",
         "Blister", "", "", "56", "stk", "6573.91", "8418"),
        ("599010", "", "Inlyta", "Pfizer AS", "Aksitinib", "Tablett", "5 mg",
         "Blister", "", "", "56", "stk", "32845.79", "41914.6"),
    ]
    for index, row in enumerate(rows, start=4):
        for column, value in enumerate(row, start=1):
            sheet.cell(index, column, value)
    path = tmp_path / "legemiddelpriser-2026-09-03.xlsx"
    workbook.save(path)
    return path


def test_parses_regulated_prices_keyed_by_item_number(register: Path) -> None:
    prices = parse_price_register(register)
    assert set(prices) == {"103854", "599010"}
    assert prices["599010"].max_aip == 32845.79
    assert prices["599010"].product_name == "Inlyta"


def test_aip_and_aup_are_kept_distinct(register: Path) -> None:
    """AIP is the pharmacy purchase ceiling; AUP adds margin and VAT.

    Hospital tenders are denominated in AIP, so confusing the two would overstate the
    price ceiling by roughly 28%.
    """
    price = parse_price_register(register)["103854"]
    assert price.max_aip == 6573.91
    assert price.max_aup == 8418.0
    assert price.max_aip != price.max_aup


def test_unreadable_register_is_not_an_error(tmp_path: Path) -> None:
    broken = tmp_path / "legemiddelpriser-broken.xlsx"
    broken.write_bytes(b"not a spreadsheet")
    assert parse_price_register(broken) == {}
